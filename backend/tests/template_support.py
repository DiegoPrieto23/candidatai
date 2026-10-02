"""Un CV de Word ficticio y su mapeo, para los tests de importación de plantillas (fase 08.3).

Los párrafos se escriben a propósito en varios fragmentos con el mismo formato, como los deja
Word al editar.
"""

import io
from pathlib import Path
from typing import Any

from docx import Document
from docx.shared import Mm, Pt, RGBColor
from PIL import Image

ACCENT = RGBColor(0x1F, 0x4E, 0x79)


def add_runs(paragraph: Any, *pieces: str, bold: bool = False, size: float | None = None) -> None:
    """Varios fragmentos con el mismo formato (lo que hace Word al corregir o revisar)."""
    for piece in pieces:
        run = paragraph.add_run(piece)
        run.bold = bold
        if size:
            run.font.size = Pt(size)


def _heading(doc: Any, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.color.rgb = ACCENT


def photo_bytes() -> bytes:
    buffer = io.BytesIO()
    Image.new("RGB", (60, 60), (150, 160, 170)).save(buffer, format="PNG")
    return buffer.getvalue()


def make_cv(path: Path, with_photo: bool = False) -> Path:
    """CV ficticio de una columna con un párrafo por cosa y los textos partidos en fragmentos."""
    doc = Document()
    name = doc.add_paragraph()
    add_runs(name, "Ana ", "Ficti", "cia Pérez", bold=True, size=20)  # 0
    if with_photo:
        name.add_run().add_picture(io.BytesIO(photo_bytes()), width=Mm(30))
    add_runs(doc.add_paragraph(), "Analista ", "de datos")  # 1
    add_runs(doc.add_paragraph(), "ana@example.com", " · ", "600 000 000")  # 2
    _heading(doc, "PERFIL")  # 3
    add_runs(doc.add_paragraph(), "Analista con dos años ", "de experiencia en informes.")  # 4
    _heading(doc, "EXPERIENCIA")  # 5
    role = doc.add_paragraph()  # 6: rol en negrita y empresa normal
    role.add_run("Analista").bold = True
    role.add_run(" · Empresa Ficticia A")
    add_runs(doc.add_paragraph(style="List Bullet"), "Automaticé ", "informes")  # 7
    add_runs(doc.add_paragraph(style="List Bullet"), "Reduje errores")  # 8
    role_b = doc.add_paragraph()  # 9
    role_b.add_run("Becaria").bold = True
    role_b.add_run(" · Empresa Ficticia B")
    add_runs(doc.add_paragraph(style="List Bullet"), "Limpié datos")  # 10
    _heading(doc, "VOLUNTARIADO")  # 11
    add_runs(doc.add_paragraph(style="List Bullet"), "Clases de ", "informática")  # 12
    _heading(doc, "IDIOMAS")  # 13
    add_runs(doc.add_paragraph(), "Español: ", "nativo")  # 14
    doc.save(str(path))
    return path


MAPPING: dict[str, Any] = {
    "paragraphs": [
        {"index": 0, "action": "variable", "text": "{{ name }}", "header": True},
        {"index": 1, "action": "variable", "text": "{{ subtitle }}", "header": True},
        {
            "index": 2,
            "action": "variable",
            "text": "{{ [contact.email, contact.phone]|select|join(' · ') }}",
            "header": True,
        },
        {
            "index": 3,
            "action": "variable",
            "text": "{{ labels.summary|upper }}",
            "start": ["if show.summary"],
        },
        {"index": 4, "action": "variable", "text": "{{ summary }}", "end": ["endif"]},
        {
            "index": 5,
            "action": "variable",
            "text": "{{ labels.experience|upper }}",
            "start": ["if show.experience"],
        },
        {
            "index": 6,
            "action": "variable",
            "runs": [
                {"text": "{{ e.role }}", "like": 0},
                {"text": " · {{ e.company }}", "like": 1},
            ],
            "start": ["for e in experience"],
        },
        {
            "index": 7,
            "action": "variable",
            "text": "{{ b }}",
            "start": ["for b in e.bullets"],
            "end": ["endfor", "endfor", "endif"],
        },
        {"index": 8, "action": "delete"},
        {"index": 9, "action": "delete"},
        {"index": 10, "action": "delete"},
        {
            "index": 11,
            "action": "variable",
            "text": "{{ x.title|upper }}",
            "start": ["if show.extra_sections", "for x in extra_sections"],
        },
        {
            "index": 12,
            "action": "variable",
            "text": "{{ i.text }}",
            "start": ["for i in x.items"],
            "end": ["endfor", "endfor", "endif"],
        },
        {
            "index": 13,
            "action": "variable",
            "text": "{{ labels.languages|upper }}",
            "start": ["if show.languages"],
        },
        {
            "index": 14,
            "action": "variable",
            "text": "{% for l in languages %}{{ l.language }}: {{ l.level }}"
            "{% if not loop.last %} · {% endif %}{% endfor %}",
            "end": ["endif"],
        },
    ]
}

"""Registro de plantillas, plantillas integradas, contexto de render y previews."""

import importlib.util
from pathlib import Path
from types import ModuleType

import pytest
from docx import Document as Docx

from app.core.config import REPO_ROOT, get_settings
from app.services.documents import context, preview, templates
from app.services.documents.engines.docx import render_docx
from tests.cv_support import demo_content, write_template
from tests.matching_support import demo_profile

SCRIPT = REPO_ROOT / "scripts" / "build_builtin_templates.py"


def load_script() -> ModuleType:
    spec = importlib.util.spec_from_file_location("build_builtin_templates", SCRIPT)
    assert spec is not None and spec.loader is not None
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def docx_with(tmp_path: Path, *paragraphs: str) -> Path:
    doc = Docx()
    for text in paragraphs:
        doc.add_paragraph(text)
    path = tmp_path / "t.docx"
    doc.save(str(path))
    return path


# --- Registro -------------------------------------------------------------------------------


def test_builtin_ats_simple_is_valid() -> None:
    template = templates.validate("ats-simple")
    assert template.meta.origin == "builtin" and template.meta.ats_friendly
    assert template.path == REPO_ROOT / "templates" / "ats-simple"
    assert (template.path / "preview.pdf").is_file()


def test_invalid_yaml_is_reported_not_raised(tmp_path: Path) -> None:
    user = get_settings().user_templates_dir
    write_template(user, "buena")
    bad = user / "mala"
    bad.mkdir()
    (bad / "template.yaml").write_text("id: mala\nengine: pdf\n", encoding="utf-8")
    found, invalid = templates.list_templates(builtin_dir=tmp_path / "vacio")
    assert [t.id for t in found] == ["buena"]
    assert [i.id for i in invalid] == ["mala"]
    assert "engine" in invalid[0].reason
    with pytest.raises(templates.TemplateNotFoundError, match="no es válida"):
        templates.get("mala", builtin_dir=tmp_path / "vacio")


def test_id_must_match_folder_and_engine_file_must_exist(tmp_path: Path) -> None:
    folder = write_template(tmp_path, "carpeta")
    (folder / "template.yaml").write_text(
        (folder / "template.yaml").read_text(encoding="utf-8").replace("id: carpeta", "id: otra"),
        encoding="utf-8",
    )
    (folder / "template.docx").unlink()
    _found, invalid = templates.list_templates(builtin_dir=tmp_path, user_dir=tmp_path / "x")
    assert "Debe coincidir con la carpeta" in invalid[0].reason
    assert "template.docx" in invalid[0].reason


def test_broken_user_template_does_not_fall_back_to_builtin(tmp_path: Path) -> None:
    builtin = tmp_path / "builtin"
    write_template(builtin, "ats-simple", origin="builtin")
    user = tmp_path / "user"
    (user / "ats-simple").mkdir(parents=True)
    (user / "ats-simple" / "template.yaml").write_text("id: ats-simple\n", encoding="utf-8")
    found, invalid = templates.list_templates(builtin_dir=builtin, user_dir=user)
    assert found == [] and [i.id for i in invalid] == ["ats-simple"]


def test_unreadable_user_template_is_listed_as_broken(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Una carpeta que creó el usuario del sandbox de un agente (Codex en Windows) sin permisos
    para el tuyo sale como rota; no tumba la lista entera."""
    builtin = tmp_path / "builtin"
    write_template(builtin, "ats-simple", origin="builtin")
    user = tmp_path / "user"
    (user / "codex").mkdir(parents=True)
    real_is_file = Path.is_file

    def is_file(self: Path) -> bool:
        if self.parent.name == "codex":
            raise PermissionError(13, "Acceso denegado", str(self))
        return real_is_file(self)

    real_load = templates.load_meta

    def load_meta(folder: Path) -> templates.TemplateMeta:
        if folder.name == "codex":
            raise PermissionError(13, "Acceso denegado", str(folder / "template.yaml"))
        return real_load(folder)

    monkeypatch.setattr(Path, "is_file", is_file)
    monkeypatch.setattr(templates, "load_meta", load_meta)
    found, invalid = templates.list_templates(builtin_dir=builtin, user_dir=user)
    assert [t.id for t in found] == ["ats-simple"]
    assert [i.id for i in invalid] == ["codex"] and "Sin permiso" in invalid[0].reason


def test_validate_rejects_variables_outside_contract(tmp_path: Path) -> None:
    docx = docx_with(tmp_path, "{{ name }}", "{{ salario_deseado }}", "{{ contact.email }}")
    write_template(get_settings().user_templates_dir, "rara", docx=docx)
    with pytest.raises(templates.TemplateInvalidError, match="salario_deseado"):
        templates.validate("rara")


def test_validate_rejects_broken_marker(tmp_path: Path) -> None:
    docx = docx_with(tmp_path, "{{ name ")
    write_template(get_settings().user_templates_dir, "rota", docx=docx)
    with pytest.raises(templates.TemplateInvalidError, match="No se puede leer"):
        templates.validate("rota")


def test_template_hash_ignores_previews(tmp_path: Path) -> None:
    folder = write_template(tmp_path, "h", docx=b"uno")
    template = templates.Template(meta=templates.load_meta(folder), path=folder)
    before = template.file_hash()
    (folder / "preview.pdf").write_bytes(b"x")
    (folder / "source.docx").write_bytes(b"y")
    assert template.file_hash() == before
    (folder / "template.docx").write_bytes(b"dos")
    assert template.file_hash() != before


# --- Plantilla integrada --------------------------------------------------------------------


BUILTIN_IDS = ("ats-simple", "clasica-foto", "carta-simple")
LETTER_ONLY = {"carta-simple": "letter.docx"}


def test_build_script_is_reproducible_and_committed_copy_is_current(tmp_path: Path) -> None:
    script = load_script()
    first, second = tmp_path / "a", tmp_path / "b"
    assert [p.name for p in script.build_all(first)] == list(BUILTIN_IDS)
    script.build_all(second)
    for template_id in BUILTIN_IDS:
        files = ("letter.docx", "template.yaml")  # las tres docx traen carta (08c)
        if template_id not in LETTER_ONLY:
            files = ("template.docx", *files)
        for name in files:
            built = (first / template_id / name).read_bytes()
            assert built == (second / template_id / name).read_bytes()
            committed = (REPO_ROOT / "templates" / template_id / name).read_bytes()
            assert built == committed, f"Regenera las plantillas: {SCRIPT.name} ({template_id})"
        if template_id not in LETTER_ONLY:
            assert (REPO_ROOT / "templates" / template_id / "preview.png").is_file()
    photo = script.build_demo_photo(tmp_path / "photo.png")
    # Los píxeles, no los bytes: el zlib de cada sistema comprime distinto (fase 26.2).
    from PIL import Image

    with Image.open(photo) as built, Image.open(REPO_ROOT / "demo" / "photo.png") as committed:
        assert built.size == committed.size and built.mode == committed.mode
        assert built.tobytes() == committed.tobytes()


def _all_paragraphs(doc: object) -> list[object]:
    paragraphs = list(doc.paragraphs)  # type: ignore[attr-defined]
    for table in doc.tables:  # type: ignore[attr-defined]
        for row in table.rows:
            for cell in row.cells:
                paragraphs.extend(cell.paragraphs)
    return paragraphs


@pytest.mark.parametrize("template_id", BUILTIN_IDS)
def test_builtin_markers_are_in_single_runs(template_id: str) -> None:
    """Ningún marcador partido entre fragmentos: si no, docxtpl no lo reconoce."""
    templates.validate(template_id, kind=None)
    engine_file = LETTER_ONLY.get(template_id, "template.docx")
    doc = Docx(str(REPO_ROOT / "templates" / template_id / engine_file))
    for paragraph in _all_paragraphs(doc):
        for run in paragraph.runs:  # type: ignore[attr-defined]
            assert run.text.count("{{") == run.text.count("}}"), run.text
            assert run.text.count("{%") == run.text.count("%}"), run.text


def test_clasica_foto_meta() -> None:
    meta = templates.get("clasica-foto").meta
    assert meta.photo.supported and meta.photo.shape == "circle"
    assert not meta.ats_friendly and meta.layout.columns == 1
    assert meta.options.accent_color is not None
    assert len(meta.options.accent_color.choices) == 4


def test_render_ats_simple_with_demo_profile(tmp_path: Path) -> None:
    template = templates.get("ats-simple")
    ctx = context.build_context(demo_content(), demo_profile(), ["certifications"])
    out = render_docx(template, ctx, tmp_path / "cv.docx", "Letter")
    doc = Docx(str(out))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "{{" not in text and "{%" not in text
    assert "Lucía Varela Otero" in text
    assert "Data Engineer — Norte Consulting (ficticia), A Coruña" in text
    assert "Lenguajes: Python, SQL" in text
    assert "Español: Nativo  ·  Inglés: B2" in text
    assert "CERTIFICACIONES" not in text  # sección oculta
    assert round(doc.sections[0].page_width.mm) == 216  # Letter


def test_render_escapes_xml_special_characters(tmp_path: Path) -> None:
    content = demo_content(summary="Pipelines <ETL> & modelos «R&D» para retail.")
    ctx = context.build_context(content, demo_profile(), [])
    out = render_docx(templates.get("ats-simple"), ctx, tmp_path / "cv.docx", "A4")
    assert "Pipelines <ETL> & modelos «R&D»" in "\n".join(p.text for p in Docx(str(out)).paragraphs)


DEMO_LINKS = {
    "https://linkedin.com/in/lucia-varela-demo",
    "https://github.com/lucia-varela-demo",
    "https://github.com/lucia-varela-demo/tfg-anomalias",
}


@pytest.mark.parametrize("template_id", ["ats-simple", "clasica-foto"])
def test_render_docx_turns_profile_links_into_hyperlinks(tmp_path: Path, template_id: str) -> None:
    ctx = context.build_context(demo_content(), demo_profile(), [])
    out = render_docx(templates.get(template_id), ctx, tmp_path / "cv.docx", "A4")
    doc = Docx(str(out))
    targets = {
        rel.target_ref for rel in doc.part.rels.values() if rel.reltype.endswith("/hyperlink")
    }
    assert targets == DEMO_LINKS
    # El texto se ve igual que sin enlace: lo que cambia es que se puede pulsar.
    text = " ".join(p.text for p in _all_paragraphs(doc))  # type: ignore[attr-defined]
    assert "github.com/lucia-varela-demo/tfg-anomalias" in text
    assert "https://" not in text


# --- Contexto -------------------------------------------------------------------------------


def test_context_project_url_and_links() -> None:
    ctx = context.build_context(demo_content(), demo_profile(), [])
    assert ctx["projects"][0]["url"] == "github.com/lucia-varela-demo/tfg-anomalias"
    assert set(context.links(ctx).values()) == DEMO_LINKS
    assert context.links({"contact": {"linkedin": "", "github": ""}, "projects": []}) == {}


@pytest.mark.parametrize(
    ("text", "expected"),
    [
        ("github.com/x", "https://github.com/x"),
        ("https://x.dev/a", "https://x.dev/a"),
        ("HTTP://x.dev", "HTTP://x.dev"),
    ],
)
def test_href_adds_scheme_only_when_missing(text: str, expected: str) -> None:
    assert context.href(text) == expected


def test_split_links_prefers_the_longest_match() -> None:
    targets = {
        "github.com/a": "https://github.com/a",
        "github.com/a/repo": "https://github.com/a/repo",
    }
    assert context.split_links("Ver github.com/a/repo y github.com/a.", targets) == [
        ("Ver ", None),
        ("github.com/a/repo", "https://github.com/a/repo"),
        (" y ", None),
        ("github.com/a", "https://github.com/a"),
        (".", None),
    ]
    assert context.split_links("nada", targets) == [("nada", None)]


@pytest.mark.parametrize(
    ("value", "language", "expected"),
    [
        ("2024-03", "es", "mar. 2024"),
        ("2024-09", "es", "sept. 2024"),
        ("2024-03", "en", "Mar 2024"),
        ("present", "es", "actualidad"),
        ("present", "en", "Present"),
        (None, "es", ""),
    ],
)
def test_format_month(value: str | None, language: str, expected: str) -> None:
    assert context.format_month(value, language) == expected  # type: ignore[arg-type]


def test_context_labels_and_show_in_english() -> None:
    ctx = context.build_context(
        demo_content(language="en", projects=[]), demo_profile(), ["languages"]
    )
    assert ctx["labels"]["experience"] == "Experience"
    assert ctx["experience"][0]["end"] == "Present"
    assert ctx["show"]["projects"] is False  # sin contenido
    assert ctx["show"]["languages"] is False  # oculta
    assert ctx["show"]["experience"] is True
    assert ctx["education"][0]["detail"] == ""
    assert ctx["skills"][0].items == ["Python", "SQL"]


# --- Preview --------------------------------------------------------------------------------


def test_preview_content_respects_capacity() -> None:
    template = templates.get("ats-simple")
    content = preview_content_with(template, bullets=3, summary_chars=80)
    assert sum(len(e.bullets) for e in content.experience) == 3
    assert [len(e.bullets) for e in content.experience] == [2, 1]  # por turnos
    assert len(content.summary) <= 81
    assert all(s.items for s in content.skills)


def preview_content_with(
    template: templates.Template, bullets: int, summary_chars: int
) -> preview.CvContent:
    meta = template.meta.model_copy(
        update={
            "capacity": template.meta.capacity.model_copy(
                update={"experience_bullets": bullets, "summary_chars": summary_chars}
            )
        }
    )
    return preview.preview_content(
        demo_profile(), templates.Template(meta=meta, path=template.path)
    )


def test_pdf_to_png_renders_first_page(tmp_path: Path) -> None:
    from PIL import Image

    from tests.cv_support import make_pdf

    pdf_file = tmp_path / "p.pdf"
    pdf_file.write_bytes(make_pdf([["Página uno"], ["Página dos"]]))
    png = preview.pdf_to_png(pdf_file, tmp_path / "out" / "p.png", width_px=400)
    with Image.open(png) as image:
        width, height = image.size
    assert abs(width - 400) <= 1 and abs(height - 400 * 842 / 595) <= 1  # pdfium redondea

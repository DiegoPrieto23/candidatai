"""Diseño de la carta a juego con el CV (fase 08c.2): integradas, html, `build-letter` y hash."""

import shutil
import zipfile
from collections.abc import Callable
from pathlib import Path

import pytest
import yaml
from docx import Document as Docx
from docx.shared import Pt
from sqlalchemy.orm import Session

from app.core.config import REPO_ROOT, get_settings
from app.models import Job
from app.schemas.cover_letter import CoverLetterContent
from app.schemas.cover_letter import CoverLetterRenderOptions as Options
from app.schemas.cv import CvRenderOptions
from app.services import events, extractions, profile
from app.services.documents import authoring, pdf, service, templates
from app.services.documents import cover_letter as letter
from app.services.documents.importer import TemplateImportError
from app.services.documents.importer import letter as letter_builder
from app.services.documents.importer.build import TemplateExistsError
from tests.cv_support import LETTER, demo_content, fake_converter, requires_chromium
from tests.matching_support import demo_extraction

DEMO_PHOTO = REPO_ROOT / "demo" / "photo.png"


@pytest.fixture(autouse=True)
def converter(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(pdf, "get_converter", lambda: fake_converter())


@pytest.fixture
def approved_job(session: Session, make_job: Callable[..., Job]) -> Job:
    profile.import_profile(session, profile.EXAMPLE_PATH)
    job = make_job("Data Engineer")
    job.location_raw = "Santiago de Compostela"
    events.record(session, job.id, "found", "system")
    extractions.save_extraction(session, job.id, demo_extraction())
    events.record(session, job.id, "approved", "user")
    session.flush()
    return job


def with_photo() -> None:
    data_dir = get_settings().data_dir
    data_dir.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(DEMO_PHOTO, data_dir / "photo.png")


def document_xml(path: Path) -> str:
    with zipfile.ZipFile(path) as z:
        return z.read("word/document.xml").decode("utf-8")


def media(path: Path) -> list[str]:
    with zipfile.ZipFile(path) as z:
        return [n for n in z.namelist() if n.startswith("word/media/")]


def save_letter(session: Session, job: Job) -> letter.SaveLetterResult:
    return letter.save_cover_letter(session, job.id, CoverLetterContent.model_validate(LETTER))


# --- Integradas ---------------------------------------------------------------------------


@pytest.mark.parametrize("template_id", ["ats-simple", "clasica-foto", "moderna-2col"])
def test_builtins_have_a_valid_letter(template_id: str) -> None:
    template = templates.validate(template_id, kind="letter")
    assert template.letter_file is not None and template.letter_file.is_file()
    assert template.meta.kind == "cv"  # siguen siendo plantillas de CV
    authoring.validate(template_id)  # CV y carta renderizan con el perfil demo


def test_letter_matches_clasica_foto_cv_with_photo_and_color(
    session: Session, approved_job: Job
) -> None:
    with_photo()
    cv = service.save_content(session, approved_job.id, demo_content())
    green = service.render(
        session, cv.document_id, CvRenderOptions(template_id="clasica-foto", accent_color="#2E5E4E")
    )
    assert green.options.photo
    service.select_render(session, green.render_id)

    render = save_letter(session, approved_job).render
    assert render.template_id == "clasica-foto" and not render.options.letter_fallback
    assert render.options.template_source == "cv" and render.warnings == []
    assert render.options.photo and render.options.accent_color == "#2E5E4E"
    docx = get_settings().data_dir / (render.docx_path or "")
    assert media(docx), "la foto va en el DOCX de la carta"
    xml = document_xml(docx)
    assert "2E5E4E" in xml and "Lucía Varela Otero" in xml

    # Sin foto en el CV, la carta tampoco la lleva.
    no_photo = service.render(
        session, cv.document_id, CvRenderOptions(template_id="clasica-foto", photo="exclude")
    )
    service.select_render(session, no_photo.render_id)
    plain = letter.render_cover_letter(session, render.document_id, Options())
    assert not plain.options.photo
    assert not media(get_settings().data_dir / (plain.docx_path or ""))


@requires_chromium
def test_moderna_2col_letter_renders_one_page(session: Session, approved_job: Job) -> None:
    doc = save_letter(session, approved_job).document_id
    render = letter.render_cover_letter(session, doc, Options(template_id="moderna-2col"))
    assert render.template_id == "moderna-2col" and render.docx_path is None
    assert render.pages == 1 and render.warnings == []


def test_letter_files_do_not_change_the_cv_hash(tmp_path: Path) -> None:
    folder = tmp_path / "ats-simple"
    shutil.copytree(templates.BUILTIN_TEMPLATES_DIR / "ats-simple", folder)
    template = templates.Template(meta=templates.load_meta(folder), path=folder)
    cv_hash, letter_hash = template.file_hash(), template.file_hash("letter")
    (folder / "letter.docx").write_bytes(b"otra carta")
    assert template.file_hash() == cv_hash and template.file_hash("letter") != letter_hash
    # El bloque `letter:` de template.yaml tampoco cuenta para el CV.
    meta = folder / "template.yaml"
    meta.write_bytes(templates.without_letter_block(meta.read_bytes()))
    assert template.file_hash() == cv_hash
    (folder / "template.docx").write_bytes(b"otro cv")
    assert template.file_hash() != cv_hash


# --- build-letter -------------------------------------------------------------------------


def fictitious_word(folder: Path) -> Path:
    """Plantilla docx FICTICIA con la cabecera en una tabla (foto | texto) y un cuerpo."""
    doc = Docx()
    table = doc.add_table(rows=1, cols=2)
    photo_cell, text_cell = table.rows[0].cells
    photo_cell.paragraphs[0].add_run("{{ photo }}")
    text_cell.paragraphs[0].add_run("{{ name }}").bold = True
    text_cell.add_paragraph("{{ subtitle }}")
    text_cell.add_paragraph("{{ contact.email }} · {{ contact.phone }}")
    doc.add_paragraph("{%p if show.summary %}")
    doc.add_paragraph("{{ labels.summary|upper }}")
    body = doc.add_paragraph().add_run("{{ summary }}")
    body.font.size = Pt(12)
    body.font.name = "Georgia"
    doc.add_paragraph("{%p endif %}")
    doc.add_paragraph("{%p if show.experience %}")
    doc.add_paragraph("{%p for e in experience %}")
    doc.add_paragraph("{{ e.role }} · {{ e.company }}")
    doc.add_paragraph("{%p endfor %}")
    doc.add_paragraph("{%p endif %}")
    folder.mkdir(parents=True)
    doc.save(str(folder / "template.docx"))
    meta = {
        "id": folder.name,
        "name": "Ficticia",
        "engine": "docx",
        "origin": "user",
        "layout": {"columns": 1, "uses_tables": True, "header_paragraphs": [0, 1, 2, 3]},
        "ats_friendly": False,
        "photo": {"supported": True, "shape": "circle", "size_mm": 25},
        "sections": ["summary", "experience"],
        "fonts": ["Georgia"],
    }
    (folder / "template.yaml").write_text(yaml.safe_dump(meta), encoding="utf-8")
    return folder


def test_build_letter_from_header_in_table(session: Session, approved_job: Job) -> None:
    folder = fictitious_word(get_settings().user_templates_dir / "ficticia")
    cv_hash = templates.get("ficticia").file_hash()

    result = authoring.build_letter("ficticia")
    assert result.header_blocks == 1 and result.body_style and result.warnings == []
    assert result.path == "templates/ficticia/letter.docx"
    template = templates.validate("ficticia", kind="letter")
    assert template.meta.letter is not None and template.meta.letter.file == "letter.docx"
    assert template.file_hash() == cv_hash  # el CV no cambia

    xml = document_xml(folder / "letter.docx")
    assert "{{ name }}" in xml and "{{ photo }}" in xml and "<w:tbl>" in xml
    assert "{{ paragraph }}" in xml and "{{ sign_off }}" in xml
    assert "summary" not in xml and "experience" not in xml
    assert "Georgia" in xml  # el cuerpo copia el formato del primer párrafo del CV

    # Y renderiza como carta de esa plantilla, con la foto del CV.
    with_photo()
    cv = service.save_content(session, approved_job.id, demo_content())
    chosen = service.render(session, cv.document_id, CvRenderOptions(template_id="ficticia"))
    service.select_render(session, chosen.render_id)
    render = save_letter(session, approved_job).render
    assert render.template_id == "ficticia" and render.options.photo
    assert media(get_settings().data_dir / (render.docx_path or ""))


def test_build_letter_errors() -> None:
    folder = fictitious_word(get_settings().user_templates_dir / "ficticia")
    authoring.build_letter("ficticia")
    with pytest.raises(TemplateExistsError, match="--force"):
        authoring.build_letter("ficticia")
    assert authoring.build_letter("ficticia", force=True).header_blocks == 1

    meta = yaml.safe_load((folder / "template.yaml").read_text(encoding="utf-8"))
    meta["layout"]["header_paragraphs"] = []
    meta.pop("letter")
    (folder / "template.yaml").write_text(yaml.safe_dump(meta), encoding="utf-8")
    with pytest.raises(TemplateImportError, match="no tiene cabecera"):
        authoring.build_letter("ficticia")

    with pytest.raises(TemplateImportError, match="integrada"):
        authoring.build_letter("clasica-foto")


def test_header_mixed_with_body_is_rejected() -> None:
    """Cabecera y cuerpo en la misma tabla (dos columnas): no se puede separar."""
    doc = Docx()
    left, right = doc.add_table(rows=1, cols=2).rows[0].cells
    left.paragraphs[0].add_run("{{ name }}")
    right.paragraphs[0].add_run("{{ summary }}")
    blocks = list(doc.element.body)[:-1]  # sin `sectPr`
    assert letter_builder.header_cut(blocks) == 0


def test_header_cut_with_nested_controls() -> None:
    doc = Docx()
    for text in ("{%p if photo %}", "{%p if subtitle %}", "{{ subtitle }}", "{%p elif name %}",
                 "{{ name }}", "{%p endif %}", "{%p endif %}", "{{ summary }}"):  # fmt: skip
        doc.add_paragraph(text)
    assert letter_builder.header_cut(list(doc.element.body)[:-1]) == 7


def test_header_cut_keeps_closing_controls() -> None:
    doc = Docx()
    for text in ("{%p if photo %}", "{{ photo }}", "{%p else %}", "{{ name }}", "{%p endif %}",
                 "{%p if show.summary %}", "{{ summary }}", "{%p endif %}"):  # fmt: skip
        doc.add_paragraph(text)
    blocks = list(doc.element.body)[:-1]
    assert letter_builder.header_cut(blocks) == 5


# --- Preview y validación -----------------------------------------------------------------


def test_letter_preview_goes_to_cache() -> None:
    result = authoring.preview_template(None, "clasica-foto", "demo", letter=True)
    assert result.pdf_path.startswith("cache/previews/clasica-foto-letter-demo-")
    assert result.pages == 1 and result.max_pages == 1
    fictitious_word(get_settings().user_templates_dir / "sin-carta")
    with pytest.raises(templates.TemplateNotFoundError, match="diseño de carta"):
        authoring.preview_template(None, "sin-carta", "demo", letter=True)


def test_declared_letter_must_exist_and_respect_contract() -> None:
    folder = fictitious_word(get_settings().user_templates_dir / "rota")
    meta_file = folder / "template.yaml"
    meta_file.write_text(
        meta_file.read_text(encoding="utf-8") + "letter:\n  file: letter.docx\n", encoding="utf-8"
    )
    found, invalid = templates.list_templates()
    assert "rota" not in [t.id for t in found] and "rota" in [i.id for i in invalid]
    doc = Docx()
    doc.add_paragraph("{{ name }} {{ summary }}")
    doc.save(str(folder / "letter.docx"))
    with pytest.raises(templates.TemplateInvalidError, match="summary"):
        templates.validate("rota", kind="letter")


def test_letter_css_is_scanned_too() -> None:
    folder = get_settings().user_templates_dir / "moderna-mia"
    shutil.copytree(templates.BUILTIN_TEMPLATES_DIR / "moderna-2col", folder)
    meta = folder / "template.yaml"
    meta.write_text(
        meta.read_text(encoding="utf-8").replace("id: moderna-2col", "id: moderna-mia"),
        encoding="utf-8",
    )
    templates.validate("moderna-mia", kind="letter")
    (folder / "letter.css").write_text(
        '@import url("https://fonts.example.com/x.css");', encoding="utf-8"
    )
    with pytest.raises(templates.TemplateInvalidError, match="letter.css"):
        templates.validate("moderna-mia", kind="letter")


def _children(element: object, ns_tag: str) -> list[str]:
    return [c.tag.split("}")[1] for c in element.findall(ns_tag)]  # type: ignore[attr-defined]


def test_letter_body_respects_ooxml_order() -> None:
    """Word exige el orden de los hijos de `pPr` y `rPr` (LibreOffice lo tolera)."""
    from docx.oxml.ns import qn

    doc = Docx()
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run("x")
    run.font.name, run.italic, run.font.size = "Georgia", True, Pt(12)
    style = (p._p.pPr, run._r.rPr)
    date, *_rest, name = letter_builder.letter_body(style)
    ppr = [c.tag.split("}")[1] for c in date.find(qn("w:pPr"))]
    assert ppr.index("pStyle") < ppr.index("spacing") < ppr.index("jc")
    rpr = [c.tag.split("}")[1] for c in name.find(qn("w:r")).find(qn("w:rPr"))]
    assert rpr.index("rFonts") < rpr.index("b") and "i" not in rpr  # sin la cursiva del CV
    assert "numPr" not in ppr


@pytest.mark.parametrize(
    "case",
    [
        b"letter:\n  file: letter.docx\n",
        b"letter:\n\n  file: letter.docx\n",
        b"letter: {file: letter.docx}\n",
    ],
)
def test_without_letter_block(case: bytes) -> None:
    base = b"id: x\nname: y\n"
    assert templates.without_letter_block(base + case) == base
    crlf = (base + case).replace(b"\n", b"\r\n")
    assert templates.without_letter_block(crlf) == base.replace(b"\n", b"\r\n")
    middle = b"id: x\n" + case + b"# fin\nname: y\n"
    assert templates.without_letter_block(middle) == b"id: x\n# fin\nname: y\n"


def test_letter_html_with_script_is_invalid() -> None:
    folder = get_settings().user_templates_dir / "moderna-js"
    shutil.copytree(templates.BUILTIN_TEMPLATES_DIR / "moderna-2col", folder)
    meta = folder / "template.yaml"
    meta.write_text(
        meta.read_text(encoding="utf-8").replace("id: moderna-2col", "id: moderna-js"),
        encoding="utf-8",
    )
    letter_html = folder / "letter.html"
    letter_html.write_text(
        letter_html.read_text(encoding="utf-8").replace("</body>", "<script>x()</script></body>"),
        encoding="utf-8",
    )
    with pytest.raises(templates.TemplateInvalidError, match="letter.html"):
        templates.validate("moderna-js", kind="letter")

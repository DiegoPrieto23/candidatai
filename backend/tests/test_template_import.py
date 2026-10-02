"""Importar un CV de Word como plantilla (fase 08.3): `inspect`, mapeo y `build-docx`.

Los Word de prueba son ficticios y se crean con python-docx (`tests/template_support.py`).
"""

import io
import json
from pathlib import Path
from typing import Any

import pytest
import yaml
from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Mm

from app.core.config import get_settings
from app.services.documents import pdf, preview, templates
from app.services.documents.engines.docx import render_docx
from app.services.documents.importer import TemplateImportError, build, docx_xml, inspect, paths
from tests.cv_support import fake_converter
from tests.template_support import MAPPING, add_runs, make_cv, photo_bytes


@pytest.fixture(autouse=True)
def converter(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(pdf, "get_converter", lambda: fake_converter())


@pytest.fixture
def uploads(isolated_env: Path) -> Path:
    folder = paths.uploads_dir()
    folder.mkdir(parents=True)
    return folder


def _render_text(template_id: str) -> tuple[str, Any]:
    template = templates.validate(template_id)
    out = get_settings().tmp_dir / "render.docx"
    render_docx(template, preview.demo_context(), out, template.meta.page.size)
    rendered = Document(str(out))
    return "\n".join(p.text for p in rendered.paragraphs), rendered


# --- inspect --------------------------------------------------------------------------------


def test_paragraphs_merge_runs_with_same_format(uploads: Path) -> None:
    doc = Document(str(make_cv(uploads / "cv.docx")))
    paragraphs = docx_xml.paragraphs(doc)
    assert [r.text for r in paragraphs[0].runs] == ["Ana Ficticia Pérez"]
    assert [(r.text, r.bold) for r in paragraphs[6].runs] == [
        ("Analista", True),
        (" · Empresa Ficticia A", False),
    ]
    assert paragraphs[7].is_list and not paragraphs[6].is_list


def test_inspect_sections_in_order_and_capacity(uploads: Path) -> None:
    report = inspect.inspect_docx(make_cv(uploads / "cv.docx"))
    assert [(s.section, s.title) for s in report.sections] == [
        ("summary", "PERFIL"),
        ("experience", "EXPERIENCIA"),
        ("extra", "VOLUNTARIADO"),
        ("languages", "IDIOMAS"),
    ]
    capacity = report.capacity
    assert capacity.summary_chars == len("Analista con dos años de experiencia en informes.")
    assert (capacity.experience_bullets, capacity.bullets_per_experience) == (3, [2, 1])
    assert capacity.skill_groups is None  # el CV no tiene sección de skills
    assert capacity.pages == 1
    assert report.ats.ats_friendly and report.tables == [] and report.images == []
    assert "#1F4E79" in report.colors
    # La fuente que se ve (la del tema), no las de estilos que el CV no usa.
    assert report.fonts == ["Cambria"]


def test_inspect_detects_tables_columns_and_image(uploads: Path) -> None:
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section._sectPr.find(qn("w:cols")).set(qn("w:num"), "2")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).paragraphs[0].add_run().add_picture(io.BytesIO(photo_bytes()), width=Mm(30))
    table.cell(0, 1).paragraphs[0].add_run("Ana Ficticia")
    doc.add_paragraph("EXPERIENCIA")
    doc.sections[0].header.paragraphs[0].text = "ana@example.com"
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    path = uploads / "dos-columnas.docx"
    doc.save(str(path))

    report = inspect.inspect_docx(path)
    assert report.page[0].columns == 2 and report.page[0].size == "A4"
    assert len(report.tables) == 1 and report.tables[0].columns == 2
    image = report.images[0]
    assert image.photo_candidate and not image.floating and image.width_mm == pytest.approx(30)
    by_index = {p.index: p for p in report.paragraphs}
    assert by_index[image.paragraph].location == "table"
    assert (by_index[image.paragraph].row, by_index[image.paragraph].cell) == (0, 0)
    assert report.header_text == ["ana@example.com"]
    assert by_index[len(report.paragraphs) - 1].location == "header"
    assert not report.ats.ats_friendly
    assert len(report.ats.reasons) == 3  # columnas, tablas y contacto en la cabecera


def test_inspect_docx_sends_other_formats_to_the_html_engine(uploads: Path) -> None:
    pdf_file = uploads / "cv.pdf"
    pdf_file.write_bytes(b"%PDF-1.4")
    with pytest.raises(TemplateImportError, match="motor html"):
        inspect.inspect_docx(pdf_file)


# --- build-docx -----------------------------------------------------------------------------


def test_build_docx_with_split_runs_renders_with_demo_profile(uploads: Path) -> None:
    source = make_cv(uploads / "mi-cv.docx")
    result = build.build_docx(source, MAPPING, "mi-cv", name="Mi CV")

    folder = get_settings().user_templates_dir / "mi-cv"
    assert result.path == "templates/mi-cv"
    assert {f.name for f in folder.iterdir()} == {
        "template.docx",
        "template.yaml",
        "source.docx",
        "source.png",
    }
    # Cada marcador, entero en un solo fragmento.
    xml = (folder / "template.docx").read_bytes()
    template_doc = Document(io.BytesIO(xml))
    for p in docx_xml.paragraphs(template_doc):
        for run in p.runs:
            assert run.text.count("{{") == run.text.count("}}")
            assert run.text.count("{%") == run.text.count("%}")
    # Conserva el formato: el rol en negrita y la empresa no.
    role = next(p for p in docx_xml.paragraphs(template_doc) if "e.role" in p.text)
    assert [(r.text, r.bold) for r in role.runs] == [
        ("{{ e.role }}", True),
        (" · {{ e.company }}", False),
    ]

    meta = yaml.safe_load((folder / "template.yaml").read_text(encoding="utf-8"))
    assert meta["sections"] == ["summary", "experience", "extra_sections", "languages"]
    assert meta["layout"]["header_paragraphs"] == [0, 1, 2]
    assert meta["capacity"]["experience_bullets"] == 3
    assert (meta["origin"], meta["source_file"], meta["name"]) == ("user", "source.docx", "Mi CV")
    assert result.sections == meta["sections"] and result.ats_friendly

    text, _ = _render_text("mi-cv")
    assert "Lucía Varela Otero" in text
    assert "Data Engineer · Norte Consulting (ficticia)" in text
    assert "Junior Data Engineer · DataGal (ficticia)" in text  # el bucle repite
    assert "Migré 40 procesos ETL" in text
    assert "VOLUNTARIADO" in text and "Mentora de análisis de datos" in text
    assert "Español: Nativo" in text
    assert "Empresa Ficticia" not in text and "{" not in text


def test_build_rejects_variable_outside_contract(uploads: Path) -> None:
    source = make_cv(uploads / "cv.docx")
    mapping = {"paragraphs": [{"index": 0, "action": "variable", "text": "{{ nombre }}"}]}
    with pytest.raises(TemplateImportError) as exc:
        build.build_docx(source, mapping, "mal")
    assert exc.value.issues[0].received == "nombre"
    assert not (get_settings().user_templates_dir / "mal").exists()


def test_build_reports_bad_attribute_and_unclosed_block_with_paragraph(uploads: Path) -> None:
    source = make_cv(uploads / "cv.docx")
    typo = {
        "paragraphs": [
            {"index": 6, "action": "variable", "text": "{{ e.rol }}",
             "start": ["for e in experience"], "end": ["endfor"]},
        ]
    }  # fmt: skip
    with pytest.raises(TemplateImportError) as exc:
        build.build_docx(source, typo, "mal")
    assert exc.value.issues[0].field == "paragraphs[index=6]"
    assert "rol" in exc.value.issues[0].reason

    unclosed = {"paragraphs": [{"index": 5, "action": "static", "start": ["if show.experience"]}]}
    with pytest.raises(TemplateImportError) as exc:
        build.build_docx(source, unclosed, "mal")
    assert "mal formados" in exc.value.issues[0].reason


@pytest.mark.parametrize(
    ("entry", "reason"),
    [
        ({"index": 99, "action": "delete"}, "No existe"),
        ({"index": 0, "action": "photo"}, "sin imagen"),
        ({"index": 0, "action": "static", "scope": "row", "start": ["if photo"],
          "end": ["endif"]}, "tabla"),
        ({"index": 6, "action": "variable", "runs": [{"text": "x", "like": 5}]}, "fuera de rango"),
    ],
)  # fmt: skip
def test_build_rejects_inconsistent_entries(
    uploads: Path, entry: dict[str, Any], reason: str
) -> None:
    source = make_cv(uploads / "cv.docx")
    with pytest.raises(TemplateImportError) as exc:
        build.build_docx(source, {"paragraphs": [entry]}, "mal")
    assert reason in exc.value.issues[0].reason


def test_mapping_schema_errors() -> None:
    from pydantic import ValidationError

    from app.services.documents.importer.mapping import DocxMapping

    bad = [
        {"index": 0, "action": "variable"},
        {"index": 1, "action": "static", "text": "x"},
        {"index": 2, "action": "delete", "start": ["if photo"]},
        {"index": 3, "action": "static", "start": ["while true"]},
    ]
    for entry in bad:
        with pytest.raises(ValidationError):
            DocxMapping.model_validate({"paragraphs": [entry]})
    with pytest.raises(ValidationError, match="repetidos"):
        DocxMapping.model_validate({"paragraphs": [{"index": 1}, {"index": 1}]})


def test_build_existing_id_needs_force(uploads: Path) -> None:
    source = make_cv(uploads / "cv.docx")
    build.build_docx(source, MAPPING, "mi-cv")
    with pytest.raises(build.TemplateExistsError):
        build.build_docx(source, MAPPING, "mi-cv")
    result = build.build_docx(source, MAPPING, "mi-cv", force=True)
    assert result.template_id == "mi-cv"


def test_build_photo_and_table_rows(uploads: Path) -> None:
    """Foto en la cabecera y una fila de tabla por experiencia (`{%tr for %}`)."""
    doc = Document()
    head = doc.add_paragraph()
    head.add_run().add_picture(io.BytesIO(photo_bytes()), width=Mm(30))  # 0
    add_runs(doc.add_paragraph(), "Ana ", "Ficticia", bold=True)  # 1
    table = doc.add_table(rows=2, cols=2)
    for r, (role, dates) in enumerate([("Analista", "2023"), ("Becaria", "2022")]):
        add_runs(table.cell(r, 0).paragraphs[0], role)  # 2 y 4
        add_runs(table.cell(r, 1).paragraphs[0], dates)  # 3 y 5
    path = uploads / "tabla.docx"
    doc.save(str(path))
    mapping = {
        "paragraphs": [
            {"index": 0, "action": "photo", "header": True},
            {"index": 1, "action": "variable", "text": "{{ name }}", "header": True},
            {"index": 2, "action": "variable", "text": "{{ e.role }}", "scope": "row",
             "start": ["for e in experience"], "end": ["endfor"]},
            {"index": 3, "action": "variable", "text": "{{ e.start }} – {{ e.end }}"},
            {"index": 4, "action": "delete"},
            {"index": 5, "action": "delete"},
        ]
    }  # fmt: skip
    result = build.build_docx(path, mapping, "tabla")
    assert result.photo_supported and result.header_paragraphs == [0, 1]
    meta = templates.get("tabla").meta
    assert (meta.photo.supported, meta.photo.size_mm, meta.layout.uses_tables) == (True, 30, True)

    _text, rendered = _render_text("tabla")
    rows = [[c.text for c in row.cells] for row in rendered.tables[0].rows]
    assert rows == [
        ["Data Engineer", "mar. 2024 – actualidad"],
        ["Junior Data Engineer", "ene. 2023 – feb. 2024"],
    ]


# --- Rutas ----------------------------------------------------------------------------------


def test_paths_must_stay_inside_data(uploads: Path, tmp_path: Path) -> None:
    outside = make_cv(tmp_path / "fuera.docx")
    with pytest.raises(paths.OutsideDataError):
        paths.source_file(outside)
    inside = make_cv(uploads / "cv.docx")
    assert paths.source_file(inside) == inside.resolve()
    tmp = get_settings().tmp_dir
    tmp.mkdir(parents=True, exist_ok=True)
    (tmp / "m.json").write_text(json.dumps(MAPPING), encoding="utf-8")
    assert paths.mapping_file(tmp / "m.json").name == "m.json"
    with pytest.raises(paths.OutsideDataError):
        paths.mapping_file(uploads / "cv.docx")
    with pytest.raises(paths.OutsideDataError):
        paths.source_file(uploads / ".." / ".." / "fuera.docx")


# --- Seguridad del mapeo y del render ---------------------------------------------------------


@pytest.mark.parametrize(
    "entry",
    [
        {"index": 0, "action": "variable", "text": "{{ ''.__class__.__mro__ }}"},
        {"index": 0, "action": "variable", "runs": [{"text": "{{ cycler.__init__ }}", "like": 0}]},
        {"index": 0, "action": "variable", "text": "{# oculto #}{{ name }}"},
        {"index": 0, "action": "static", "start": ["if x %}{% for y in z"], "end": ["endif"]},
    ],
)
def test_mapping_rejects_python_internals_and_comments(entry: dict[str, Any]) -> None:
    from pydantic import ValidationError

    from app.services.documents.importer.mapping import DocxMapping

    with pytest.raises(ValidationError):
        DocxMapping.model_validate({"paragraphs": [entry]})


def test_render_is_sandboxed(tmp_path: Path) -> None:
    """Aunque una plantilla llegue con un marcador peligroso, el render no lo ejecuta."""
    from jinja2.exceptions import SecurityError

    from tests.cv_support import write_template

    doc = Document()
    doc.add_paragraph("{{ name.upper.__self__.__class__.__mro__ }}")
    source = tmp_path / "malo.docx"
    doc.save(str(source))
    folder = write_template(get_settings().user_templates_dir, "malo", docx=source)
    template = templates.get("malo")
    assert template.path == folder
    with pytest.raises(SecurityError):
        render_docx(template, preview.demo_context(), tmp_path / "out.docx", "A4")

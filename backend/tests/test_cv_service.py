"""Servicio de CVs: guardar contenido, renderizar, verificar y llegar a `cv_ready` (fase 08.1)."""

import shutil
import zipfile
from collections.abc import Callable
from pathlib import Path

import pytest
import yaml
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.core.config import REPO_ROOT, get_settings
from app.models import Document, Event, Job
from app.schemas.cv import CvRenderOptions, CvVerification
from app.services import events, extractions, profile
from app.services.documents import checks, options, pdf, service, templates
from app.services.pipeline import Actor, InvalidTransitionError, Stage
from tests.cv_support import demo_content, demo_content_data, fake_converter, write_template
from tests.matching_support import demo_extraction, demo_profile

REAL_GET_CONVERTER = pdf.get_converter
ATS_SIMPLE_DOCX = REPO_ROOT / "templates" / "ats-simple" / "template.docx"
PASSED = CvVerification(passed=True, issues=[])


@pytest.fixture(autouse=True)
def converter(monkeypatch: pytest.MonkeyPatch) -> None:
    """Sin LibreOffice: un conversor falso que escribe el texto del DOCX en un PDF mínimo."""
    monkeypatch.setattr(pdf, "get_converter", lambda: fake_converter())


def use_converter(monkeypatch: pytest.MonkeyPatch, convert: Callable[..., None]) -> None:
    monkeypatch.setattr(pdf, "get_converter", lambda: convert)


@pytest.fixture
def approved_job(session: Session, make_job: Callable[..., Job]) -> Job:
    """Oferta aprobada con la extracción demo y el perfil demo importado."""
    profile.import_profile(session, profile.EXAMPLE_PATH)
    job = make_job("Data Engineer")
    job.location_raw = "Santiago de Compostela"
    events.record(session, job.id, "found", "system")
    extractions.save_extraction(session, job.id, demo_extraction())
    events.record(session, job.id, "approved", "user")
    session.flush()
    return job


def stage(session: Session, job: Job) -> Stage:
    return events.current_stage(session, job.id)


def event_types(session: Session, job: Job) -> list[str]:
    stmt = select(Event.type).where(Event.job_id == job.id).order_by(Event.id)
    return list(session.scalars(stmt))


def data_file(relative: str | None) -> Path:
    assert relative
    return get_settings().data_dir / relative


# --- Guardar y renderizar -------------------------------------------------------------------


def test_save_content_stores_document_and_first_render(session: Session, approved_job: Job) -> None:
    result = service.save_content(session, approved_job.id, demo_content())

    doc = session.get(Document, result.document_id)
    assert doc is not None and doc.kind == "cv" and doc.language == "es"
    render = result.render
    # El perfil demo pide `clasica-foto` por defecto; sin `data/photo.*`, va sin foto y sin aviso.
    assert render.template_id == "clasica-foto"
    assert render.options.template_source == "default_template"
    assert render.options.photo is False and render.options.accent_color == "#1F4E79"
    assert render.warnings == []
    assert render.pages == 1 and render.max_pages == 1
    assert render.selected is True
    assert (
        render.docx_path
        == f"cvs/{approved_job.id}/{doc.id}/{render.render_id}/CV-Lucia-Varela-Otero.docx"
    )
    assert data_file(render.docx_path).is_file() and data_file(render.pdf_path).is_file()
    assert render.options.offer.countries == ["ES"]
    assert "Python" in render.keyword_coverage.present
    assert event_types(session, approved_job)[-2:] == ["cv_generated", "cv_rendered"]
    assert stage(session, approved_job) == Stage.APPROVED


def test_contact_comes_from_profile_not_from_agent(session: Session, approved_job: Job) -> None:
    result = service.save_content(session, approved_job.id, demo_content())
    from docx import Document as Docx

    text = "\n".join(p.text for p in Docx(str(data_file(result.render.docx_path))).paragraphs)
    assert "+34 600 000 000" in text and "lucia.varela@example.com" in text
    assert "Lucía Varela Otero" in text
    assert "mar. 2024 – actualidad" in text
    # Idioma regional: el redactor no lo pidió y el gallego no aparece aunque lo incluyera.
    assert "Gallego" not in text


def test_regional_language_only_when_requested(session: Session, approved_job: Job) -> None:
    langs = [{"language": "Español", "level": "Nativo"}, {"language": "Gallego", "level": "Nativo"}]
    from docx import Document as Docx

    for include, expected in ((False, False), (True, True)):
        content = demo_content(languages=langs, include_regional_language=include)
        result = service.save_content(session, approved_job.id, content)
        docx_path = data_file(result.render.docx_path)
        text = "\n".join(p.text for p in Docx(str(docx_path)).paragraphs)
        assert ("Gallego" in text) is expected


def test_verification_passed_reaches_cv_ready(session: Session, approved_job: Job) -> None:
    saved = service.save_content(session, approved_job.id, demo_content())
    result = service.save_verification(session, saved.document_id, PASSED)
    assert (result.passed, result.stage, result.note) == (True, Stage.CV_READY, None)
    assert event_types(session, approved_job)[-1] == "cv_verified"


def test_verification_failed_is_rejected(session: Session, approved_job: Job) -> None:
    saved = service.save_content(session, approved_job.id, demo_content())
    issue = {
        "severity": "blocker",
        "location": "experience[0].bullets[0]",
        "problem": "La cifra no está en el perfil.",
        "suggestion": "Usa -35 %.",
    }
    verification = CvVerification.model_validate({"passed": False, "issues": [issue]})
    result = service.save_verification(session, saved.document_id, verification)
    assert (result.passed, result.stage) == (False, Stage.APPROVED)
    last = session.scalars(select(Event).order_by(Event.id.desc())).first()
    assert last is not None and (last.type, last.payload["blockers"]) == ("cv_rejected", 1)
    view = service.get_content(session, document_id=saved.document_id)
    assert view.verification == verification


def test_passed_with_blocker_is_invalid() -> None:
    with pytest.raises(ValueError, match="blocker"):
        CvVerification.model_validate(
            {"passed": True, "issues": [{"severity": "blocker", "location": "x", "problem": "y"}]}
        )


def test_only_approved_jobs_get_a_cv(
    session: Session, make_job: Callable[..., Job], approved_job: Job
) -> None:
    job = make_job("Otra")
    events.record(session, job.id, "found", "system")
    with pytest.raises(InvalidTransitionError, match="solo se generan CVs"):
        service.save_content(session, job.id, demo_content())
    assert session.scalar(select(func.count()).select_from(Document)) == 0


# --- Contenido rechazado --------------------------------------------------------------------


def _issues(exc: pytest.ExceptionInfo[checks.CvContentError]) -> list[tuple[str, str | None]]:
    return [(i.field, i.received) for i in exc.value.issues]


def test_rejects_unknown_evidence_and_saves_nothing(session: Session, approved_job: Job) -> None:
    data = demo_content_data()
    data["experience"][0]["bullets"][1]["evidence"] = ["exp-norteconsulting/a9"]
    data["projects"][0]["bullets"][0]["evidence"] = ["prj-inventado"]
    with pytest.raises(checks.CvContentError) as exc:
        service.save_content(session, approved_job.id, demo_content(**data))
    assert _issues(exc) == [
        ("experience[0].bullets[1].evidence[0]", "exp-norteconsulting/a9"),
        ("projects[0].bullets[0].evidence[0]", "prj-inventado"),
    ]
    assert session.scalar(select(func.count()).select_from(Document)) == 0
    assert "cv_generated" not in event_types(session, approved_job)


def test_rejects_unknown_ids() -> None:
    content = demo_content(
        experience=[
            {"experience_id": "exp-nada", "bullets": [{"text": "x", "evidence": ["exp-datagal"]}]}
        ],
        education_ids=["edu-nada"],
        certification_ids=["cert-nada"],
        projects=[],
    )
    with pytest.raises(checks.CvContentError) as exc:
        checks.validate_content(content, demo_profile(), demo_extraction())
    assert [f for f, _ in _issues(exc)] == [
        "experience[0].experience_id",
        "education_ids[0]",
        "certification_ids[0]",
    ]


def test_rejects_skill_not_in_profile() -> None:
    content = demo_content(skills=[{"group": "Datos", "items": ["Python", "Kubernetes"]}])
    with pytest.raises(checks.CvContentError) as exc:
        checks.validate_content(content, demo_profile(), demo_extraction())
    assert _issues(exc) == [("skills[0].items[1]", "Kubernetes")]


def test_skill_aliases_count_as_the_profile_skill() -> None:
    content = demo_content(skills=[{"group": "Lenguajes", "items": ["python", "PowerBI"]}])
    checks.validate_content(content, demo_profile(), demo_extraction())


def test_rejects_basic_skill_not_requested() -> None:
    extraction = demo_extraction(must_have_skills=["Python"], nice_to_have_skills=[])
    content = demo_content(skills=[{"group": "Datos", "items": ["conocimientos de dbt"]}])
    with pytest.raises(checks.CvContentError, match="skills") as exc:
        checks.validate_content(content, demo_profile(), extraction)
    assert "básica" in exc.value.issues[0].reason


@pytest.mark.parametrize("item", ["dbt", "conocimientos de dbt", "Conocimientos básicos de DBT"])
def test_basic_skill_allowed_when_requested(item: str) -> None:
    extraction = demo_extraction(must_have_skills=["Python"], nice_to_have_skills=["dbt"])
    content = demo_content(skills=[{"group": "Datos", "items": [item]}])
    checks.validate_content(content, demo_profile(), extraction)


def test_level_b_skill_counts_as_basic() -> None:
    data = demo_profile().model_dump(mode="json")
    data["skills"].append({"name": "Kafka", "level": "B", "category": "Datos"})
    from app.schemas.profile import CandidateProfile

    prof = CandidateProfile.model_validate(data)
    content = demo_content(skills=[{"group": "Datos", "items": ["Kafka"]}])
    with pytest.raises(checks.CvContentError):
        checks.validate_content(content, prof, demo_extraction(nice_to_have_skills=[]))
    checks.validate_content(content, prof, demo_extraction(nice_to_have_skills=["Kafka"]))


def test_basic_skill_without_extraction_is_rejected() -> None:
    content = demo_content(skills=[{"group": "Datos", "items": ["dbt"]}])
    with pytest.raises(checks.CvContentError):
        checks.validate_content(content, demo_profile(), None)


# --- Renders --------------------------------------------------------------------------------


@pytest.fixture
def no_ats_template() -> Path:
    """Una plantilla del usuario igual que `ats-simple` pero no apta para ATS."""
    return write_template(
        get_settings().user_templates_dir, "no-ats", docx=ATS_SIMPLE_DOCX, ats_friendly=False
    )


def test_second_template_does_not_touch_document(
    session: Session, approved_job: Job, no_ats_template: Path
) -> None:
    saved = service.save_content(session, approved_job.id, demo_content())
    doc = session.get(Document, saved.document_id)
    assert doc is not None
    before = (dict(doc.content), doc.created_at)

    other = service.render(session, saved.document_id, CvRenderOptions(template_id="no-ats"))
    assert other.template_id == "no-ats" and other.options.template_source == "option"
    assert other.selected is False  # sigue seleccionado el primero
    session.refresh(doc)
    assert (dict(doc.content), doc.created_at) == before
    assert session.scalar(select(func.count()).select_from(Document)) == 1
    assert [
        r.template_id for r in service.get_content(session, job_id=approved_job.id).renders
    ] == [
        "clasica-foto",
        "no-ats",
    ]


def test_not_ats_friendly_warning_with_workday_url(
    session: Session, approved_job: Job, no_ats_template: Path
) -> None:
    approved_job.url = "https://acme.wd3.myworkdayjobs.com/es-ES/careers/job/123"
    saved = service.save_content(
        session, approved_job.id, demo_content(), CvRenderOptions(template_id="no-ats")
    )
    codes = [w.code for w in saved.render.warnings]
    assert "not_ats_friendly" in codes
    assert saved.render.options.offer.ats == "workday"
    # Con `ats-simple` no hay aviso.
    again = service.render(session, saved.document_id, CvRenderOptions(template_id="ats-simple"))
    assert "not_ats_friendly" not in [w.code for w in again.warnings]


def test_two_pages_is_too_long(
    session: Session, approved_job: Job, monkeypatch: pytest.MonkeyPatch
) -> None:
    use_converter(monkeypatch, fake_converter(lines_per_page=15))
    saved = service.save_content(session, approved_job.id, demo_content())
    render = saved.render
    assert render.pages == 2
    assert "too_long" in [w.code for w in render.warnings]
    assert render.selected is False
    last = session.scalars(select(Event).order_by(Event.id.desc())).first()
    assert last is not None and last.payload["too_long"] is True

    # Verificado, pero sin render que quepa: no pasa a `cv_ready` y se explica por qué.
    result = service.save_verification(session, saved.document_id, PASSED)
    assert result.stage == Stage.APPROVED and result.note

    # Otro render que sí cabe (dos páginas permitidas) lo lleva a `cv_ready` y queda seleccionado.
    fits = service.render(session, saved.document_id, CvRenderOptions(max_pages=2))
    assert "too_long" not in [w.code for w in fits.warnings]
    assert fits.selected is True
    assert stage(session, approved_job) == Stage.CV_READY


def test_font_substituted_warning(
    session: Session, approved_job: Job, monkeypatch: pytest.MonkeyPatch
) -> None:
    use_converter(monkeypatch, fake_converter(font="Carlito"))
    saved = service.save_content(session, approved_job.id, demo_content())
    assert "font_substituted" in [w.code for w in saved.render.warnings]


def test_section_not_supported_warning(session: Session, approved_job: Job) -> None:
    write_template(
        get_settings().user_templates_dir,
        "corta",
        docx=ATS_SIMPLE_DOCX,
        sections=["summary", "experience", "skills"],
    )
    saved = service.save_content(
        session, approved_job.id, demo_content(), CvRenderOptions(template_id="corta")
    )
    warning = next(w for w in saved.render.warnings if w.code == "section_not_supported")
    assert "projects" in warning.message and "summary" not in warning.message
    # Ocultar la sección quita el aviso.
    hidden = service.render(
        session,
        saved.document_id,
        CvRenderOptions(
            template_id="corta", hidden_sections=["projects", "education", "certifications"]
        ),
    )
    warning = next(w for w in hidden.warnings if w.code == "section_not_supported")
    assert warning.message.endswith("languages.")


def test_hidden_section_is_not_rendered(session: Session, approved_job: Job) -> None:
    saved = service.save_content(
        session, approved_job.id, demo_content(), CvRenderOptions(hidden_sections=["projects"])
    )
    from docx import Document as Docx

    text = "\n".join(p.text for p in Docx(str(data_file(saved.render.docx_path))).paragraphs)
    assert "PROYECTOS" not in text and "EXPERIENCIA" in text


def test_invalid_accent_color_is_rejected(session: Session, approved_job: Job) -> None:
    saved = service.save_content(session, approved_job.id, demo_content())
    with pytest.raises(options.InvalidRenderOptionsError, match="no admite color"):
        service.render(
            session,
            saved.document_id,
            CvRenderOptions(template_id="ats-simple", accent_color="#1F4E79"),
        )
    with pytest.raises(options.InvalidRenderOptionsError, match="admite estos colores"):
        service.render(session, saved.document_id, CvRenderOptions(accent_color="#000000"))


def test_html_template_is_not_available_yet(session: Session, approved_job: Job) -> None:
    folder = write_template(get_settings().user_templates_dir, "web", engine="html")
    (folder / "template.html").write_text("<p>{{ name }}</p>", encoding="utf-8")
    (folder / "style.css").write_text("", encoding="utf-8")
    saved = service.save_content(session, approved_job.id, demo_content())
    with pytest.raises(service.EngineNotAvailableError, match="08b"):
        service.render(session, saved.document_id, CvRenderOptions(template_id="web"))


def test_failed_conversion_leaves_no_files(
    session: Session, approved_job: Job, monkeypatch: pytest.MonkeyPatch
) -> None:
    def broken(_docx: Path, _pdf: Path, _timeout: int) -> None:
        raise pdf.PdfConversionError("LibreOffice no generó el PDF.")

    use_converter(monkeypatch, broken)
    with pytest.raises(pdf.PdfConversionError):
        service.save_content(session, approved_job.id, demo_content())
    cvs = get_settings().data_dir / "cvs"
    assert not cvs.exists() or not any(p.is_file() for p in cvs.rglob("*"))


def test_no_converter_explains_what_to_install(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(pdf, "find_soffice", lambda _configured: None)
    monkeypatch.setattr(pdf, "word_converter", lambda: None)
    with pytest.raises(pdf.PdfConverterUnavailableError, match="LibreOffice"):
        REAL_GET_CONVERTER()


# --- Foto, colores y renders múltiples (08.2) -----------------------------------------------

DEMO_PHOTO = REPO_ROOT / "demo" / "photo.png"


def docx_media(relative: str | None) -> list[str]:
    with zipfile.ZipFile(data_file(relative)) as docx:
        return [n for n in docx.namelist() if n.startswith("word/media/")]


def docx_xml(relative: str | None) -> str:
    with zipfile.ZipFile(data_file(relative)) as docx:
        return docx.read("word/document.xml").decode("utf-8")


def selected_ids(session: Session, job: Job) -> list[int]:
    renders = service.get_content(session, job_id=job.id).renders
    return [r.render_id for r in renders if r.selected]


@pytest.fixture
def with_photo() -> Path:
    target = get_settings().data_dir / "photo.png"
    shutil.copyfile(DEMO_PHOTO, target)
    return target


def test_clasica_foto_with_photo_and_four_colors(
    session: Session, approved_job: Job, with_photo: Path
) -> None:
    saved = service.save_content(session, approved_job.id, demo_content())
    render = saved.render
    assert render.template_id == "clasica-foto" and render.options.photo is True
    assert render.warnings == []
    assert len(docx_media(render.docx_path)) == 1
    cache = get_settings().data_dir / "cache" / "photo"
    assert len(list(cache.glob("*-circle-*.png"))) == 1

    accent = templates.get("clasica-foto").meta.options.accent_color
    assert accent is not None and len(accent.choices) == 4
    for color in accent.choices:
        result = service.render(session, saved.document_id, CvRenderOptions(accent_color=color))
        xml = docx_xml(result.docx_path)
        assert f'w:val="{color[1:]}"' in xml
        assert not any(f'w:val="{c[1:]}"' in xml for c in accent.choices if c != color)
        assert result.options.accent_color == color
        assert len(docx_media(result.docx_path)) == 1
    assert len(list(cache.glob("*.png"))) == 1  # la foto recortada se reutiliza


def test_clasica_foto_without_photo(session: Session, approved_job: Job, with_photo: Path) -> None:
    saved = service.save_content(
        session, approved_job.id, demo_content(), CvRenderOptions(photo="exclude")
    )
    assert saved.render.options.photo is False and docx_media(saved.render.docx_path) == []
    from docx import Document as Docx

    doc = Docx(str(data_file(saved.render.docx_path)))
    assert doc.tables == []  # sin foto, la cabecera va en el cuerpo, sin tabla
    assert "Lucía Varela Otero" in "\n".join(p.text for p in doc.paragraphs)


def test_photo_requested_but_missing_or_unsupported(session: Session, approved_job: Job) -> None:
    saved = service.save_content(
        session, approved_job.id, demo_content(), CvRenderOptions(photo="include")
    )
    warning = next(w for w in saved.render.warnings if w.code == "photo_excluded")
    assert "no hay foto" in warning.message
    assert docx_media(saved.render.docx_path) == []

    shutil.copyfile(DEMO_PHOTO, get_settings().data_dir / "photo.jpg")
    ats = service.render(
        session, saved.document_id, CvRenderOptions(template_id="ats-simple", photo="include")
    )
    warning = next(w for w in ats.warnings if w.code == "photo_excluded")
    assert "no la admite" in warning.message


def test_photo_auto_follows_country_policy(
    session: Session, approved_job: Job, with_photo: Path
) -> None:
    approved_job.location_raw = "Dublin, Ireland"  # España (extracción) e Irlanda: sin foto
    saved = service.save_content(
        session, approved_job.id, demo_content(), CvRenderOptions(template_id="clasica-foto")
    )
    assert "IE" in saved.render.options.offer.countries
    assert saved.render.options.photo is False and docx_media(saved.render.docx_path) == []
    forced = service.render(
        session, saved.document_id, CvRenderOptions(template_id="clasica-foto", photo="include")
    )
    assert forced.options.photo is True and len(docx_media(forced.docx_path)) == 1


def test_switch_template_and_photo_without_touching_document(
    session: Session, approved_job: Job, with_photo: Path
) -> None:
    saved = service.save_content(session, approved_job.id, demo_content())
    doc = session.get(Document, saved.document_id)
    assert doc is not None
    before = (dict(doc.content), doc.verification, doc.created_at)
    events_before = event_types(session, approved_job)

    ats = service.render(
        session, saved.document_id, CvRenderOptions(template_id="ats-simple"), Actor.USER
    )
    no_photo = service.render(
        session, saved.document_id, CvRenderOptions(photo="exclude"), Actor.USER
    )
    session.refresh(doc)
    assert (dict(doc.content), doc.verification, doc.created_at) == before
    assert session.scalar(select(func.count()).select_from(Document)) == 1
    assert (ats.template_id, ats.options.photo) == ("ats-simple", False)
    assert (no_photo.template_id, no_photo.options.photo) == ("clasica-foto", False)
    assert len({saved.render.render_id, ats.render_id, no_photo.render_id}) == 3
    assert event_types(session, approved_job) == [*events_before, "cv_rendered", "cv_rendered"]
    last = session.scalars(select(Event).order_by(Event.id.desc())).first()
    assert last is not None and last.actor == "user"


def test_select_render_keeps_a_single_one(session: Session, approved_job: Job) -> None:
    saved = service.save_content(session, approved_job.id, demo_content())
    other = service.render(session, saved.document_id, CvRenderOptions(template_id="ats-simple"))
    assert (saved.render.selected, other.selected) == (True, False)
    stage_before, events_before = stage(session, approved_job), event_types(session, approved_job)

    assert service.select_render(session, other.render_id).selected is True
    assert selected_ids(session, approved_job) == [other.render_id]
    service.select_render(session, other.render_id)  # el mismo otra vez: nada cambia
    service.select_render(session, saved.render.render_id)
    assert selected_ids(session, approved_job) == [saved.render.render_id]
    # Sin evento ni cambio de etapa.
    assert stage(session, approved_job) == stage_before
    assert event_types(session, approved_job) == events_before

    with pytest.raises(service.RenderNotFoundError):
        service.select_render(session, 999)


def test_render_marks_template_changed(
    session: Session, approved_job: Job, no_ats_template: Path
) -> None:
    saved = service.save_content(
        session, approved_job.id, demo_content(), CvRenderOptions(template_id="no-ats")
    )
    assert saved.render.template_changed is False
    assert service.get_content(session, job_id=approved_job.id).renders[0].template_changed is False
    docx = no_ats_template / "template.docx"
    docx.write_bytes(docx.read_bytes() + b"\0")
    assert service.get_content(session, job_id=approved_job.id).renders[0].template_changed is True
    shutil.rmtree(no_ats_template)  # si ya no existe, también cuenta como cambiada
    assert service.get_content(session, job_id=approved_job.id).renders[0].template_changed is True


def test_max_pages_from_profile_preference(session: Session, approved_job: Job) -> None:
    data = profile.current(session).profile.model_dump(mode="json")
    data["preferences"]["cv"]["max_pages"] = 2
    path = get_settings().data_dir / "profile-2pag.yaml"
    path.write_text(yaml.safe_dump(data, allow_unicode=True), encoding="utf-8")
    profile.import_profile(session, path)

    saved = service.save_content(session, approved_job.id, demo_content())
    assert saved.render.max_pages == 2
    # La opción de render manda sobre el perfil.
    one = service.render(session, saved.document_id, CvRenderOptions(max_pages=1))
    assert one.max_pages == 1


# --- Ver el contenido -----------------------------------------------------------------------


def test_get_content_resolves_evidence(session: Session, approved_job: Job) -> None:
    saved = service.save_content(session, approved_job.id, demo_content())
    view = service.get_content(session, job_id=approved_job.id)
    assert view.document_id == saved.document_id
    assert view.evidence["exp-norteconsulting/a1"].startswith("Migré 40 procesos ETL")
    assert "-35 %" in view.evidence["exp-norteconsulting/a1"]
    assert list(view.evidence) == [
        "exp-norteconsulting/a1",
        "exp-norteconsulting/a2",
        "exp-datagal/a1",
        "prj-tfg",
    ]
    assert [r.render_id for r in view.renders] == [saved.render.render_id]


def test_get_content_errors(session: Session, approved_job: Job) -> None:
    with pytest.raises(service.DocumentNotFoundError, match="aún no tiene CV"):
        service.get_content(session, job_id=approved_job.id)
    with pytest.raises(service.DocumentNotFoundError):
        service.get_content(session, document_id=999)
    with pytest.raises(ValueError, match="uno de los dos"):
        service.get_content(session)


def test_list_cv_templates_includes_builtin_and_invalid() -> None:
    folder = get_settings().user_templates_dir / "rota"
    folder.mkdir(parents=True)
    (folder / "template.yaml").write_text("id: rota\n", encoding="utf-8")
    listed = service.list_cv_templates()
    ids = [t.id for t in listed.templates]
    assert "ats-simple" in ids
    ats = next(t for t in listed.templates if t.id == "ats-simple")
    assert ats.ats_friendly and not ats.photo_supported and ats.capacity["experience_bullets"]
    assert [i.id for i in listed.invalid] == ["rota"]
    assert listed.fallback_template == "ats-simple"


# --- Con LibreOffice de verdad --------------------------------------------------------------

needs_soffice = pytest.mark.skipif(
    pdf.find_soffice(None) is None, reason="LibreOffice no está instalado"
)


@needs_soffice
def test_real_render_of_ats_simple_with_demo_profile(
    session: Session, approved_job: Job, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(pdf, "get_converter", REAL_GET_CONVERTER)
    saved = service.save_content(
        session, approved_job.id, demo_content(), CvRenderOptions(template_id="ats-simple")
    )
    render = saved.render
    assert render.pages == 1
    assert render.warnings == []  # sin font_substituted
    assert {"Python", "SQL", "Snowflake"} <= set(render.keyword_coverage.present)
    info = pdf.read_pdf(data_file(render.pdf_path))
    assert "Lucía Varela Otero" in info.text
    assert not list(Path(get_settings().data_dir).glob("**/jh-soffice-*"))


@needs_soffice
def test_real_render_of_clasica_foto_with_photo(
    session: Session, approved_job: Job, monkeypatch: pytest.MonkeyPatch, with_photo: Path
) -> None:
    monkeypatch.setattr(pdf, "get_converter", REAL_GET_CONVERTER)
    saved = service.save_content(session, approved_job.id, demo_content())
    render = saved.render
    assert (render.template_id, render.options.photo, render.pages) == ("clasica-foto", True, 1)
    assert render.warnings == []  # Cambria y Calibri disponibles
    assert "Lucía Varela Otero" in pdf.read_pdf(data_file(render.pdf_path)).text


# --- Secciones extra (fase 08.3) ------------------------------------------------------------

EXTRA = [{"section_id": "voluntariado", "item_ids": ["v1"]}]


def test_rejects_unknown_extra_section_or_item() -> None:
    content = demo_content(
        extra_sections=[
            {"section_id": "premios", "item_ids": ["p1"]},
            {"section_id": "voluntariado", "item_ids": ["v1", "v9"]},
        ]
    )
    with pytest.raises(checks.CvContentError) as exc:
        checks.validate_content(content, demo_profile(), demo_extraction())
    assert _issues(exc) == [
        ("extra_sections[0].section_id", "premios"),
        ("extra_sections[1].item_ids[1]", "v9"),
    ]


def test_repeated_extra_items_are_invalid() -> None:
    with pytest.raises(ValueError, match="repite ids"):
        demo_content(extra_sections=[{"section_id": "voluntariado", "item_ids": ["v1", "v1"]}])


@pytest.mark.parametrize("template_id", ["ats-simple", "clasica-foto"])
def test_extra_sections_rendered_by_builtins(
    session: Session, approved_job: Job, template_id: str
) -> None:
    from docx import Document as Docx

    saved = service.save_content(
        session,
        approved_job.id,
        demo_content(extra_sections=EXTRA),
        CvRenderOptions(template_id=template_id),
    )
    assert not [w for w in saved.render.warnings if w.code == "section_not_supported"]
    text = "\n".join(p.text for p in Docx(str(data_file(saved.render.docx_path))).paragraphs)
    assert "VOLUNTARIADO" in text and "Mentora de análisis de datos" in text
    assert "oct. 2023" in text

    hidden = service.render(
        session,
        saved.document_id,
        CvRenderOptions(template_id=template_id, hidden_sections=["extra_sections"]),
    )
    text = "\n".join(p.text for p in Docx(str(data_file(hidden.docx_path))).paragraphs)
    assert "VOLUNTARIADO" not in text


def test_extra_section_title_in_cv_language(session: Session, approved_job: Job) -> None:
    from app.services.documents import context

    content = demo_content(language="en", extra_sections=EXTRA)
    ctx = context.build_context(content, demo_profile(), [])
    assert ctx["extra_sections"][0].title == "Volunteering"
    assert ctx["extra_sections"][0].items == [
        {"text": "Mentora de análisis de datos en un programa de iniciación para estudiantes de FP",
         "date": "Oct 2023"}
    ]  # fmt: skip
    assert ctx["show"]["extra_sections"] is True
    assert (
        context.build_context(demo_content(), demo_profile(), [])["show"]["extra_sections"] is False
    )


def test_template_without_extra_sections_warns(session: Session, approved_job: Job) -> None:
    write_template(get_settings().user_templates_dir, "sin-extra", docx=ATS_SIMPLE_DOCX)
    saved = service.save_content(
        session,
        approved_job.id,
        demo_content(extra_sections=EXTRA),
        CvRenderOptions(template_id="sin-extra"),
    )
    warning = next(w for w in saved.render.warnings if w.code == "section_not_supported")
    assert "extra_sections" in warning.message

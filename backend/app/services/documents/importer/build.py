"""`cai templates build-docx`: aplica el mapeo al Word del usuario y crea la plantilla.

Cada párrafo mapeado se reescribe con un solo fragmento (o uno por parte de `runs`) con el formato
del original, así que ningún marcador queda partido. Los bloques se insertan como párrafos o filas
de control de docxtpl (`{%p … %}`, `{%tr … %}`), que desaparecen al renderizar. El resultado se
valida renderizándolo con el perfil demo; si falla, no se deja nada a medias.
"""

import copy
import shutil
import tempfile
from pathlib import Path
from typing import Any

import yaml
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pydantic import BaseModel

from app.core.config import get_settings
from app.schemas.cv import TemplateCapacity, TemplateMeta
from app.services.documents import pdf, templates
from app.services.documents.importer import (
    ImportIssue,
    TemplateImportError,
    docx_xml,
    inspect,
    mapping,
)

SOURCE_DOCX = "source.docx"
SOURCE_PDF = "source.pdf"
SOURCE_PNG = "source.png"
_SYMBOL_FONTS = {"symbol", "wingdings", "wingdings 2", "wingdings 3", "webdings"}


class TemplateExistsError(ValueError):
    pass


class BuildResult(BaseModel):
    template_id: str
    path: str  # relativa a data/
    sections: list[str]
    header_paragraphs: list[int]
    capacity: dict[str, int]
    max_pages: int
    photo_supported: bool
    ats_friendly: bool
    ats_reasons: list[str]
    source_preview: bool
    warnings: list[str]


# --- XML ------------------------------------------------------------------------------------


def _run(rpr: Any, text: str) -> Any:
    """Un fragmento con el formato `rpr` y el texto (tabuladores y saltos como en Word)."""
    r = OxmlElement("w:r")
    if rpr is not None:
        r.append(copy.deepcopy(rpr))
    for i, line in enumerate(text.split("\n")):
        if i:
            r.append(OxmlElement("w:br"))
        for j, chunk in enumerate(line.split("\t")):
            if j:
                r.append(OxmlElement("w:tab"))
            if chunk:
                t = OxmlElement("w:t")
                t.set(qn("xml:space"), "preserve")
                t.text = chunk
                r.append(t)
    return r


def _rpr(run: docx_xml.Run | None) -> Any:
    return run.elements[0].find(qn("w:rPr")) if run is not None else None


def _rewrite(p: docx_xml.Paragraph, entry: mapping.ParagraphMap) -> None:
    parts: list[tuple[Any, str]]
    if entry.runs is not None:
        parts = [(_rpr(p.runs[part.like] if p.runs else None), part.text) for part in entry.runs]
    else:
        first = next((r for r in p.runs if r.text.strip()), p.runs[0] if p.runs else None)
        parts = [(_rpr(first), entry.text or "")]
    for child in list(p.element):
        if child.tag != qn("w:pPr"):
            p.element.remove(child)
    for rpr, text in parts:
        p.element.append(_run(rpr, text))


def _replace_photo(p: docx_xml.Paragraph) -> None:
    image_runs = [r for r in p.element.iter(docx_xml.W_R) if docx_xml.has_image(r)]
    first = image_runs[0]
    first.addprevious(_run(first.find(qn("w:rPr")), "{{ photo }}"))
    for r in image_runs:
        r.getparent().remove(r)


def _control_paragraph(text: str) -> Any:
    p = OxmlElement("w:p")
    p.append(_run(None, text))
    return p


def _control_row(tr: Any, text: str) -> Any:
    """Copia de la fila con las celdas vacías y la etiqueta en la primera."""
    row = copy.deepcopy(tr)
    for tc in row.findall(docx_xml.W_TC):
        for child in list(tc):
            if child.tag != qn("w:tcPr"):
                tc.remove(child)
        tc.append(OxmlElement("w:p"))
    first_p = row.find(f"{docx_xml.W_TC}/{docx_xml.W_P}")
    first_p.append(_run(None, text))
    return row


def _row_of(p: Any) -> Any:
    return next(p.iterancestors(docx_xml.W_TR))


def _insert_blocks(p: docx_xml.Paragraph, entry: mapping.ParagraphMap) -> None:
    row_scope = entry.scope == "row"
    anchor = _row_of(p.element) if row_scope else p.element
    prefix = "{%tr " if row_scope else "{%p "
    for tag in entry.start:
        text = f"{prefix}{tag.strip()} %}}"
        anchor.addprevious(_control_row(anchor, text) if row_scope else _control_paragraph(text))
    last = anchor
    for tag in entry.end:
        text = f"{prefix}{tag} %}}"
        new = _control_row(anchor, text) if row_scope else _control_paragraph(text)
        last.addnext(new)
        last = new


def _delete(paragraphs: list[docx_xml.Paragraph], deleted: set[int]) -> None:
    rows: dict[int, tuple[Any, list[int]]] = {}
    for p in paragraphs:
        tr = next(p.element.iterancestors(docx_xml.W_TR), None)
        if tr is not None:
            rows.setdefault(id(tr), (tr, []))[1].append(p.index)
    for p in paragraphs:
        if p.index in deleted:
            p.element.getparent().remove(p.element)
    # Una fila cuyos párrafos se han borrado todos sobra (el bucle repite la primera).
    for tr, indexes in rows.values():
        if all(i in deleted for i in indexes) and tr.getparent() is not None:
            tbl = tr.getparent()
            tbl.remove(tr)
            if tbl.find(docx_xml.W_TR) is None and tbl.getparent() is not None:
                tbl.getparent().remove(tbl)
    # Word exige al menos un párrafo en cada celda y cuadro de texto.
    for container_tag in (docx_xml.W_TC, docx_xml.W_TXBX):
        for p in paragraphs:
            for container in p.element.iterancestors(container_tag):
                if container.find(docx_xml.W_P) is None and container.find(docx_xml.W_TBL) is None:
                    container.append(OxmlElement("w:p"))


def _drop_fallbacks(roots: list[Any]) -> None:
    """Las copias de compatibilidad de los cuadros de texto conservarían el texto original."""
    for root in roots:
        for fallback in list(root.iter(docx_xml.FALLBACK)):
            fallback.getparent().remove(fallback)


# --- Metadatos ------------------------------------------------------------------------------


def _capacity(measured: inspect.MeasuredCapacity) -> TemplateCapacity:
    defaults = TemplateCapacity()
    return TemplateCapacity(
        summary_chars=measured.summary_chars or defaults.summary_chars,
        experience_bullets=measured.experience_bullets or defaults.experience_bullets,
        skill_groups=measured.skill_groups or defaults.skill_groups,
    )


def _meta(
    template_id: str,
    name: str | None,
    report: inspect.DocxInspection,
    sections: list[str],
    header: list[int],
    photo_paragraph: int | None,
    pages: int | None,
) -> dict[str, Any]:
    photo = next((i for i in report.images if i.paragraph == photo_paragraph), None)
    page = report.page[0] if report.page else None
    layout_tables = [t for t in report.tables if t.borderless or t.columns > 1]
    columns = max([p.columns for p in report.page] + [t.columns for t in layout_tables] + [1])
    fonts = [f for f in report.fonts if f.casefold() not in _SYMBOL_FONTS][:2]
    data = {
        "id": template_id,
        "name": name or template_id.replace("-", " ").capitalize(),
        "engine": "docx",
        "origin": "user",
        "source_file": SOURCE_DOCX,
        "description": f"Creada desde {Path(report.file).name}.",
        "tags": [],
        "layout": {
            "columns": min(columns, 3),
            "uses_tables": bool(report.tables or report.text_boxes),
            "header_paragraphs": header,
        },
        "ats_friendly": report.ats.ats_friendly,
        "photo": {
            "supported": photo_paragraph is not None,
            "shape": photo.shape if photo else "circle",
            "size_mm": min(max(photo.width_mm, 10), 60) if photo else 30,
        },
        "page": {
            "size": page.size if page and page.size != "other" else "A4",
            "max_pages": min(max(pages or 1, 1), 3),
        },
        "capacity": _capacity(report.capacity).model_dump(),
        "sections": sections,
        "fonts": fonts,
        "languages": ["es", "en"],
    }
    TemplateMeta.model_validate(data)  # lo que escribimos tiene que ser válido
    return data


# --- Entrada --------------------------------------------------------------------------------


def _source_pdf(source: Path, folder: Path) -> tuple[int | None, bool, str | None]:
    """Páginas del original y su primera página como `source.png` (para «Original | Recreación»)."""
    from app.services.documents.preview import pdf_to_png

    workdir = Path(tempfile.mkdtemp(prefix="cai-source-"))
    try:
        copy_docx = workdir / SOURCE_DOCX
        shutil.copyfile(source, copy_docx)
        pdf.convert(copy_docx, workdir / SOURCE_PDF)
        pages = pdf.read_pdf(workdir / SOURCE_PDF).pages
        pdf_to_png(workdir / SOURCE_PDF, folder / SOURCE_PNG)
        return pages, True, None
    except pdf.PdfConversionError as exc:
        return None, False, f"Sin vista del original ni páginas: {exc}"
    finally:
        shutil.rmtree(workdir, ignore_errors=True)


def build_docx(
    source: Path,
    mapping_data: Any,
    template_id: str,
    name: str | None = None,
    force: bool = False,
) -> BuildResult:
    from app.services.documents.preview import check_renders

    # El id con el mismo formato que en `template.yaml`, antes de hacer nada.
    TemplateMeta.model_validate(
        {"id": template_id, "name": "x", "engine": "docx", "origin": "user",
         "ats_friendly": False, "sections": ["summary"]}
    )  # fmt: skip
    parsed = mapping.DocxMapping.model_validate(mapping_data)
    report = inspect.inspect_docx(source, with_pages=False)
    doc = inspect.open_docx(source)
    paragraphs = docx_xml.paragraphs(doc)
    mapping.check(paragraphs, parsed)
    sections = mapping.used_sections(paragraphs, parsed)
    if not sections:
        issue = ImportIssue(
            field="mapping", received=None, reason="La plantilla no usa ninguna sección del CV."
        )
        raise TemplateImportError("El mapeo no usa ninguna sección del contrato.", [issue])

    user_dir = get_settings().user_templates_dir
    folder = user_dir / template_id
    if folder.exists() and not force:
        raise TemplateExistsError(
            f"Ya existe la plantilla «{template_id}» en {folder}. Usa otro id o `--force` para "
            "sustituirla."
        )

    warnings: list[str] = []
    entries = parsed.by_index()
    photo_paragraph = None
    deleted = set()
    for p in paragraphs:
        entry = entries.get(p.index)
        if entry is None:
            continue
        if entry.action == "variable":
            uses_photo = "photo" in entry.content
            if p.images and not uses_photo:
                warnings.append(f"Párrafo {p.index}: su imagen se pierde al reescribirlo.")
            if uses_photo:
                photo_paragraph = p.index
            _rewrite(p, entry)
        elif entry.action == "photo":
            photo_paragraph = p.index
            _replace_photo(p)
        elif entry.action == "delete":
            deleted.add(p.index)
        _insert_blocks(p, entry)
    _delete(paragraphs, deleted)
    _drop_fallbacks([doc.element.body, *(el for _k, el in docx_xml.header_footer_parts(doc))])
    floating = [i for i in report.images if i.paragraph == photo_paragraph and i.floating]
    if floating:
        warnings.append(
            "La foto del original flotaba sobre el texto; en la plantilla va en línea, en su "
            "párrafo. Revisa la preview."
        )
    if (templates.BUILTIN_TEMPLATES_DIR / template_id / templates.META_FILE).is_file():
        warnings.append(f"«{template_id}» es también una integrada: la tuya la sustituye.")
    header = sorted(p.index for p in parsed.paragraphs if p.header)

    settings = get_settings()
    settings.tmp_dir.mkdir(parents=True, exist_ok=True)
    staging = Path(tempfile.mkdtemp(prefix="cai-build-", dir=settings.tmp_dir))
    try:
        target = staging / template_id
        target.mkdir()
        doc.save(str(target / "template.docx"))
        shutil.copyfile(source, target / SOURCE_DOCX)
        pages, has_png, note = _source_pdf(source, target)
        if note:
            warnings.append(note)
        meta = _meta(template_id, name, report, sections, header, photo_paragraph, pages)
        (target / "template.yaml").write_text(
            yaml.safe_dump(meta, allow_unicode=True, sort_keys=False), encoding="utf-8"
        )
        template = templates.Template(meta=templates.load_meta(target), path=target)
        try:
            check_renders(template)
        except templates.TemplateInvalidError as exc:
            issues = [
                ImportIssue(field=i.field, received=None, reason=i.reason) for i in exc.issues
            ]
            message = f"La plantilla resultante no renderiza: {exc}"
            raise TemplateImportError(message, issues) from exc
        if folder.exists():
            shutil.rmtree(folder)
        user_dir.mkdir(parents=True, exist_ok=True)
        shutil.move(str(target), str(folder))
    finally:
        shutil.rmtree(staging, ignore_errors=True)

    return BuildResult(
        template_id=template_id,
        path=folder.relative_to(settings.data_dir).as_posix(),
        sections=sections,
        header_paragraphs=header,
        capacity=meta["capacity"],
        max_pages=meta["page"]["max_pages"],
        photo_supported=meta["photo"]["supported"],
        ats_friendly=report.ats.ats_friendly,
        ats_reasons=report.ats.reasons,
        source_preview=has_png,
        warnings=warnings,
    )
